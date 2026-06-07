"""
tools/file_reader.py — 文档类文件读取（Word / PDF / TXT）

返回结构化的文档内容，供 /api/upload 预览和 Agent 上下文注入使用。
"""

from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional


@dataclass
class TableResult:
    headers: list          # 表头列名列表
    rows: list             # 数据行列表（每行为列值列表）


@dataclass
class DocumentResult:
    ok: bool
    file_type: str = ""       # "docx" | "pdf" | "txt"
    text: str = ""            # 提取的纯文本
    word_count: int = 0
    page_count: int = 0
    tables: list = field(default_factory=list)   # list[TableResult]
    error: str = ""


def read_document(file_path: str, max_chars: int = 5000) -> DocumentResult:
    """
    读取文档类文件，提取文本内容。

    支持：
      .docx — python-docx
      .pdf  — pypdf
      .txt  — 直接读取（自适应编码）
    """
    p = Path(file_path)
    ext = p.suffix.lower()

    if ext == '.docx':
        return _read_docx(file_path, max_chars)
    elif ext == '.pdf':
        return _read_pdf(file_path, max_chars)
    elif ext == '.txt':
        return _read_txt(file_path, max_chars)
    else:
        return DocumentResult(ok=False, error=f"不支持的文件类型：{ext}")


def _read_docx(file_path: str, max_chars: int) -> DocumentResult:
    try:
        from docx import Document
        doc = Document(file_path)
    except Exception as e:
        return DocumentResult(ok=False, file_type="docx", error=f"无法读取 Word 文件：{str(e)[:200]}")

    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    text = "\n".join(paragraphs)

    tables = []
    for tbl in doc.tables:
        all_rows = [[cell.text.strip() for cell in row.cells] for row in tbl.rows]
        if all_rows:
            headers = all_rows[0]
            data_rows = all_rows[1:]
            tables.append(TableResult(headers=headers, rows=data_rows))

    return DocumentResult(
        ok=True,
        file_type="docx",
        text=text[:max_chars],
        word_count=len(text.split()),
        page_count=0,
        tables=tables,
    )


def _read_pdf(file_path: str, max_chars: int) -> DocumentResult:
    try:
        from pypdf import PdfReader
    except ImportError:
        try:
            from PyPDF2 import PdfReader
        except ImportError:
            return DocumentResult(ok=False, file_type="pdf",
                                  error="缺少 PDF 解析库，请运行 pip install pypdf")

    try:
        reader = PdfReader(file_path)
        pages = []
        for page in reader.pages:
            pages.append(page.extract_text() or "")
        text = "\n".join(pages)
        return DocumentResult(
            ok=True,
            file_type="pdf",
            text=text[:max_chars],
            word_count=len(text.split()),
            page_count=len(reader.pages),
            tables=[],
        )
    except Exception as e:
        return DocumentResult(ok=False, file_type="pdf", error=f"无法读取 PDF 文件：{str(e)[:200]}")


def _read_txt(file_path: str, max_chars: int) -> DocumentResult:
    from tools.data_loader import detect_encoding
    try:
        enc = detect_encoding(file_path)
        text = Path(file_path).read_text(encoding=enc, errors='replace')
        return DocumentResult(
            ok=True,
            file_type="txt",
            text=text[:max_chars],
            word_count=len(text.split()),
            page_count=1,
            tables=[],
        )
    except FileNotFoundError:
        return DocumentResult(ok=False, file_type="txt", error=f"文件不存在：{file_path}")
    except Exception as e:
        return DocumentResult(ok=False, file_type="txt", error=f"无法读取文本文件：{str(e)[:200]}")
