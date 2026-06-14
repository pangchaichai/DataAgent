"""
tools/report_builder.py — 报告渲染与 Word 导出
"""

import re
import time
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent.parent
TEMPLATES_DIR = BASE_DIR / 'templates' / 'reports'
OUTPUTS_DIR = BASE_DIR / 'data' / 'outputs'

_TEMPLATE_REGISTRY = [
    {"name": "concentration_report", "label": "主体/单券集中度报告", "path": "concentration_report.md.j2"},
    {"name": "nav_report",           "label": "净值运作报告",         "path": "nav_report.md.j2"},
]


@dataclass
class ReportResult:
    ok: bool
    markdown: str = ""
    error: str = ""


@dataclass
class WordResult:
    ok: bool
    word_path: str = ""
    error: str = ""


# ═══════════════════════════════════════════════════════════════
#  Public API
# ═══════════════════════════════════════════════════════════════

def list_templates() -> list:
    """返回所有可用报告模板列表（每项含 name/label/path）"""
    return [t.copy() for t in _TEMPLATE_REGISTRY]


def render_report(template_name: str, data: dict) -> ReportResult:
    """
    使用模板名称渲染报告，返回 ReportResult。
    data 中若未提供 report_date，自动注入当日日期。
    """
    entry = next((t for t in _TEMPLATE_REGISTRY if t["name"] == template_name), None)
    if not entry:
        return ReportResult(ok=False, error=f"模板不存在：{template_name}")

    tmpl_path = TEMPLATES_DIR / entry["path"]
    if not tmpl_path.exists():
        return ReportResult(ok=False, error=f"模板文件不存在：{tmpl_path}")

    ctx = dict(data)
    ctx.setdefault("report_date", datetime.now().strftime("%Y-%m-%d"))

    try:
        from jinja2 import Environment, FileSystemLoader, StrictUndefined
        env = Environment(
            loader=FileSystemLoader(str(BASE_DIR / 'templates')),
            undefined=StrictUndefined,
        )
        template = env.get_template('reports/' + entry["path"])
        md = template.render(**ctx)
        return ReportResult(ok=True, markdown=md)
    except Exception as e:
        return ReportResult(ok=False, error=f"模板渲染失败：{e}")


def export_word(markdown_text: str, output_path: str) -> WordResult:
    """
    将 Markdown 文本导出为 Word .docx 文件。
    返回 WordResult(ok, word_path, error)。
    """
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
    except ImportError:
        return WordResult(ok=False, error="缺少 python-docx，请运行 pip install python-docx")

    doc = Document()
    style = doc.styles['Normal']
    style.font.name = '宋体'
    style.font.size = Pt(11)

    lines = markdown_text.split('\n')
    i = 0
    while i < len(lines):
        line = lines[i]

        if line.startswith('### '):
            doc.add_heading(line[4:].strip(), level=3)
        elif line.startswith('## '):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith('# '):
            doc.add_heading(line[2:].strip(), level=1)
        elif re.match(r'^[-=]{3,}$', line.strip()):
            doc.add_paragraph('─' * 30)
        elif line.strip().startswith('|'):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                table_lines.append(lines[i])
                i += 1
            _add_table(doc, table_lines)
            continue
        elif re.match(r'^\s*[-*]\s+', line):
            text = re.sub(r'^\s*[-*]\s+', '', line)
            p = doc.add_paragraph(style='List Bullet')
            _add_runs(p, text)
        elif re.match(r'^\s*\d+\.\s+', line):
            text = re.sub(r'^\s*\d+\.\s+', '', line)
            p = doc.add_paragraph(style='List Number')
            _add_runs(p, text)
        elif line.startswith('> '):
            p = doc.add_paragraph(line[2:].strip())
            p.style = doc.styles['Normal']
            p.paragraph_format.left_indent = Pt(20)
            if p.runs:
                p.runs[0].font.color.rgb = RGBColor(0x5C, 0x63, 0x70)
        elif not line.strip():
            doc.add_paragraph('')
        else:
            p = doc.add_paragraph()
            _add_runs(p, line)

        i += 1

    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    OUTPUTS_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return WordResult(ok=True, word_path=str(out))


def save_report(content: str, filename: str) -> str:
    """保存报告 Markdown 到 data/outputs/ 并返回文件路径"""
    OUTPUTS_DIR.mkdir(parents=True, exist_ok=True)
    out_path = OUTPUTS_DIR / filename
    out_path.write_text(content, encoding='utf-8')
    return str(out_path)


def export_report_word(content: str, report_name: str) -> dict:
    """
    将报告 Markdown 字符串导出为 Word 文件，保存到 data/outputs/。
    返回 {"ok": True, "filename": str, "path": str} 或 {"ok": False, "error": str}
    """
    ts = time.strftime("%Y%m%d_%H%M%S")
    filename = f"{report_name}_{ts}.docx"
    out_path = str(OUTPUTS_DIR / filename)
    result = export_word(content, out_path)
    if result.ok:
        return {"ok": True, "filename": filename, "path": result.word_path}
    return {"ok": False, "error": result.error}


# ═══════════════════════════════════════════════════════════════
#  Internal helpers
# ═══════════════════════════════════════════════════════════════

def _add_runs(paragraph, text: str):
    parts = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('*') and part.endswith('*'):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        elif part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Courier New'
        elif part:
            paragraph.add_run(part)


def _add_table(doc, lines: list):
    rows_data = []
    for line in lines:
        if re.match(r'^\|[\s\-:|]+\|', line):
            continue
        cells = [c.strip() for c in line.strip('|').split('|')]
        rows_data.append(cells)

    if not rows_data:
        return

    max_cols = max(len(r) for r in rows_data)
    table = doc.add_table(rows=len(rows_data), cols=max_cols)
    table.style = 'Table Grid'

    for r_idx, row_data in enumerate(rows_data):
        row = table.rows[r_idx]
        for c_idx, cell_text in enumerate(row_data):
            if c_idx < max_cols:
                cell = row.cells[c_idx]
                cell.text = cell_text
                if r_idx == 0:
                    for run in cell.paragraphs[0].runs:
                        run.bold = True
