"""Excel 预处理器：处理合并单元格、标题行、多层表头、多Sheet拼接。

场景 A：标题行（单合并格横跨 ≥80% 列宽）→ 自动跳过
场景 B：双层表头（横跨+竖跨混合）→ 填充合并值 → 逐列拼接去重
场景 C：多Sheet同构 → 逐Sheet处理 → 纵向合并
"""

from __future__ import annotations

import os
from dataclasses import dataclass, field

import pandas as pd

TITLE_ROW_THRESHOLD = 0.8


@dataclass
class SheetInfo:
    name: str
    row_count: int
    col_count: int
    title_rows_skipped: int
    header_levels: int
    columns: list[str]


@dataclass
class PreprocessResult:
    df: pd.DataFrame
    sheet_info: list[SheetInfo]
    warnings: list[str] = field(default_factory=list)
    sheets_concatenated: bool = False


def preprocess_excel(file_path: str) -> PreprocessResult:
    """主入口：自动检测标题行、多层表头、多Sheet，返回干净 DataFrame。"""
    ext = os.path.splitext(file_path)[1].lower()
    if ext not in ('.xlsx', '.xls'):
        raise ValueError(f"不支持的文件格式：{ext}")

    raw_sheets = _read_raw_sheets(file_path, ext)

    if not raw_sheets:
        return PreprocessResult(df=pd.DataFrame(), sheet_info=[],
                                warnings=["文件无有效工作表"])

    processed = []
    for name, rows, merged in raw_sheets:
        result = _process_single_sheet(name, rows, merged)
        if result is not None:
            processed.append(result)

    if not processed:
        return PreprocessResult(df=pd.DataFrame(), sheet_info=[],
                                warnings=["所有工作表均无有效数据"])

    if len(processed) == 1:
        df, info = processed[0]
        return PreprocessResult(df=df, sheet_info=[info])

    all_info = [info for _, info in processed]
    if _check_column_compatibility(all_info):
        combined = pd.concat([df for df, _ in processed], ignore_index=True)
        total = sum(i.row_count for i in all_info)
        return PreprocessResult(
            df=combined, sheet_info=all_info, sheets_concatenated=True,
            warnings=[f"多Sheet列名一致，已纵向合并（共 {total} 行）"])

    df, info = processed[0]
    return PreprocessResult(
        df=df, sheet_info=all_info, sheets_concatenated=False,
        warnings=[f"此文件包含 {len(all_info)} 个工作表（结构不同），"
                  f"仅导入首张工作表「{info.name}」"])


# ── 内部函数 ───────────────────────────────────────────────


def _read_raw_sheets(file_path, ext):
    """读取所有工作表，返回 [(name, 2D_values, merged_ranges_0based)]。"""
    sheets = []
    if ext == '.xlsx':
        import openpyxl
        wb = openpyxl.load_workbook(file_path, data_only=True)
        for ws in wb.worksheets:
            rows = [list(row) for row in ws.iter_rows(values_only=True)]
            merged = [(mr.min_row - 1, mr.max_row - 1,
                       mr.min_col - 1, mr.max_col - 1)
                      for mr in ws.merged_cells.ranges]
            if rows:
                sheets.append((ws.title, rows, merged))
        wb.close()
    else:
        import xlrd
        wb = xlrd.open_workbook(file_path, formatting_info=False)
        for i in range(wb.nsheets):
            ws = wb.sheet_by_index(i)
            rows = [[ws.cell_value(r, c) for c in range(ws.ncols)]
                    for r in range(ws.nrows)]
            merged = [(rlo, rhi - 1, clo, chi - 1)
                      for rlo, rhi, clo, chi in ws.merged_cells]
            if rows:
                sheets.append((ws.name, rows, merged))
    return sheets


def _fill_merged_cells(rows, merged_ranges):
    """将合并区域左上角的值填充到所有被合并的格子。"""
    for min_r, max_r, min_c, max_c in merged_ranges:
        if min_r >= len(rows):
            continue
        val = rows[min_r][min_c] if min_c < len(rows[min_r]) else None
        for r in range(min_r, min(max_r + 1, len(rows))):
            for c in range(min_c, min(max_c + 1, len(rows[r]))):
                rows[r][c] = val


def _detect_title_rows(rows, total_cols, merged_ranges):
    """从第1行起连续检测标题行，返回跳过的行数。

    两种判定方式（任一满足即跳过）：
    1. 合并格元数据：该行有且仅有1个合并格，且跨度 ≥ 总列数的80%
    2. 启发式：该行仅有 ≤1 个非空单元格（适用于无合并信息的 .xls）
    """
    skip = 0
    for r_idx, row in enumerate(rows):
        wide_merges = [
            1 for min_r, _, min_c, max_c in merged_ranges
            if min_r == r_idx
            and (max_c - min_c + 1) >= total_cols * TITLE_ROW_THRESHOLD
        ]
        if len(wide_merges) == 1:
            skip += 1
            continue

        non_empty = sum(1 for v in row
                        if v is not None and str(v).strip() != '')
        if non_empty <= 1 and total_cols >= 3:
            skip += 1
            continue

        break
    return skip


def _detect_header_levels(rows, start_row, merged_ranges):
    """从 start_row 起，检测参与合并格的连续行数作为表头层数。"""
    if not merged_ranges:
        return 1
    levels = 0
    for r_idx in range(start_row, len(rows)):
        participates = any(min_r <= r_idx <= max_r
                          for min_r, max_r, _, _ in merged_ranges)
        if participates:
            levels += 1
        else:
            break
    return max(levels, 1)


def _build_column_names(header_rows):
    """逐列拼接各层表头值，去除相邻重复，用 _ 连接。"""
    if not header_rows:
        return []
    num_cols = max(len(row) for row in header_rows)
    columns = []
    for c in range(num_cols):
        parts: list[str] = []
        for row in header_rows:
            val = row[c] if c < len(row) else None
            s = str(val).strip() if val is not None else ''
            if s and (not parts or parts[-1] != s):
                parts.append(s)
        columns.append('_'.join(parts) if parts else f'列{c + 1}')

    seen: dict[str, int] = {}
    result: list[str] = []
    for name in columns:
        if name in seen:
            seen[name] += 1
            result.append(f'{name}_{seen[name]}')
        else:
            seen[name] = 0
            result.append(name)
    return result


def _to_str(val):
    """将单元格值转为字符串，None/NaN → 空字符串。"""
    if val is None:
        return ''
    if isinstance(val, float):
        if val != val:
            return ''
        if val == int(val) and abs(val) < 1e15:
            return str(int(val))
        return str(val)
    s = str(val)
    return '' if s in ('None', 'nan', 'NaN') else s


def _process_single_sheet(name, rows, merged_ranges):
    """处理单个工作表：填充→跳过标题→检测表头→构建 DataFrame。"""
    if not rows:
        return None
    total_cols = max(len(row) for row in rows)
    if total_cols == 0:
        return None

    has_data = any(
        any(v is not None and str(v).strip() != '' for v in row)
        for row in rows
    )
    if not has_data:
        return None

    for row in rows:
        while len(row) < total_cols:
            row.append(None)

    _fill_merged_cells(rows, merged_ranges)

    title_skip = _detect_title_rows(rows, total_cols, merged_ranges)
    header_start = title_skip
    if header_start >= len(rows):
        return None

    header_levels = _detect_header_levels(rows, header_start, merged_ranges)
    header_rows = rows[header_start:header_start + header_levels]
    columns = _build_column_names(header_rows)

    data_start = header_start + header_levels
    data_rows = rows[data_start:]

    if not data_rows:
        df = pd.DataFrame(columns=columns)
    else:
        str_rows = [[_to_str(v) for v in row[:len(columns)]] for row in data_rows]
        df = pd.DataFrame(str_rows, columns=columns)

    info = SheetInfo(name=name, row_count=len(data_rows), col_count=len(columns),
                     title_rows_skipped=title_skip, header_levels=header_levels,
                     columns=columns)
    return df, info


def _check_column_compatibility(sheet_infos):
    """检查所有工作表的列名是否完全一致。"""
    if len(sheet_infos) <= 1:
        return True
    first = sheet_infos[0].columns
    return all(info.columns == first for info in sheet_infos[1:])
