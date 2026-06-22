"""Tests for tools/excel_preprocessor.py — Excel 预处理器。"""

import os
import tempfile

import openpyxl
import pytest

from tools.excel_preprocessor import (
    PreprocessResult,
    _build_column_names,
    _to_str,
    preprocess_excel,
)


@pytest.fixture
def tmp_dir():
    with tempfile.TemporaryDirectory() as d:
        yield d


def _create_xlsx(path, sheets_data):
    """创建测试 xlsx 文件。

    sheets_data: [(sheet_name, rows_2d, merges_list)]
    merges_list: ["A1:C1", ...] openpyxl 格式
    """
    wb = openpyxl.Workbook()
    for i, (name, rows, merges) in enumerate(sheets_data):
        ws = wb.active if i == 0 else wb.create_sheet(name)
        if i == 0:
            ws.title = name
        for r_idx, row in enumerate(rows, 1):
            for c_idx, val in enumerate(row, 1):
                ws.cell(row=r_idx, column=c_idx, value=val)
        for merge_range in merges:
            ws.merge_cells(merge_range)
    wb.save(path)


# ── 场景 A：标题行检测 ────────────────────────────────────


class TestTitleRowDetection:

    def test_title_row_skipped(self, tmp_dir):
        path = os.path.join(tmp_dir, "title.xlsx")
        _create_xlsx(path, [("Sheet1", [
            ["数据报表标题", None, None, None, None],
            ["列A", "列B", "列C", "列D", "列E"],
            ["v1", "v2", "v3", "v4", "v5"],
        ], ["A1:E1"])])

        result = preprocess_excel(path)
        assert result.sheet_info[0].title_rows_skipped == 1
        assert result.df.columns.tolist() == ["列A", "列B", "列C", "列D", "列E"]
        assert len(result.df) == 1

    def test_title_row_not_skipped_below_threshold(self, tmp_dir):
        path = os.path.join(tmp_dir, "no_title.xlsx")
        row1 = ["标题", None, None, None, None, None, None, "X", "Y", "Z"]
        row2 = list("ABCDEFGHIJ")
        row3 = list("1234567890")
        _create_xlsx(path, [("Sheet1", [row1, row2, row3], ["A1:G1"])])

        result = preprocess_excel(path)
        assert result.sheet_info[0].title_rows_skipped == 0

    def test_no_title_no_merges(self, tmp_dir):
        path = os.path.join(tmp_dir, "simple.xlsx")
        _create_xlsx(path, [("Sheet1", [
            ["列A", "列B", "列C"],
            ["v1", "v2", "v3"],
        ], [])])

        result = preprocess_excel(path)
        assert result.sheet_info[0].title_rows_skipped == 0
        assert result.df.columns.tolist() == ["列A", "列B", "列C"]
        assert len(result.df) == 1

    def test_multiple_title_rows(self, tmp_dir):
        path = os.path.join(tmp_dir, "multi_title.xlsx")
        _create_xlsx(path, [("Sheet1", [
            ["大标题", None, None, None],
            ["副标题", None, None, None],
            ["A", "B", "C", "D"],
            ["1", "2", "3", "4"],
        ], ["A1:D1", "A2:D2"])])

        result = preprocess_excel(path)
        assert result.sheet_info[0].title_rows_skipped == 2
        assert result.df.columns.tolist() == ["A", "B", "C", "D"]


# ── 场景 B：双层表头 ──────────────────────────────────────


class TestMultiLevelHeader:

    def test_two_level_header(self, tmp_dir):
        path = os.path.join(tmp_dir, "header2.xlsx")
        _create_xlsx(path, [("Sheet1", [
            ["资产名称", "总资产", None, "净资产", None],
            [None, "成本", "市值", "成本", "市值"],
            ["AAA", "100", "200", "300", "400"],
        ], ["A1:A2", "B1:C1", "D1:E1"])])

        result = preprocess_excel(path)
        assert result.sheet_info[0].header_levels == 2
        expected = ["资产名称", "总资产_成本", "总资产_市值",
                    "净资产_成本", "净资产_市值"]
        assert result.df.columns.tolist() == expected
        assert len(result.df) == 1
        assert result.df.iloc[0]["资产名称"] == "AAA"

    def test_vertical_merge_dedup(self, tmp_dir):
        path = os.path.join(tmp_dir, "vmerge.xlsx")
        _create_xlsx(path, [("Sheet1", [
            ["名称", "偏离金额"],
            [None, None],
            ["X", "100"],
        ], ["A1:A2", "B1:B2"])])

        result = preprocess_excel(path)
        assert result.df.columns.tolist() == ["名称", "偏离金额"]
        assert result.sheet_info[0].header_levels == 2

    def test_three_col_horizontal_merge(self, tmp_dir):
        path = os.path.join(tmp_dir, "h3merge.xlsx")
        _create_xlsx(path, [("Sheet1", [
            ["净值波动%", None, None, "偏离金额"],
            ["较上一日", "较上月", "较年初", None],
            ["-0.5", "1.2", "-3.0", "100"],
        ], ["A1:C1", "D1:D2"])])

        result = preprocess_excel(path)
        expected = ["净值波动%_较上一日", "净值波动%_较上月",
                    "净值波动%_较年初", "偏离金额"]
        assert result.df.columns.tolist() == expected


# ── 场景 C：多Sheet拼接 ──────────────────────────────────


class TestMultiSheet:

    def test_same_columns_concatenated(self, tmp_dir):
        path = os.path.join(tmp_dir, "multi.xlsx")
        _create_xlsx(path, [
            ("S1", [["A", "B"], ["1", "2"], ["3", "4"]], []),
            ("S2", [["A", "B"], ["5", "6"]], []),
            ("S3", [["A", "B"], ["7", "8"], ["9", "10"]], []),
        ])

        result = preprocess_excel(path)
        assert result.sheets_concatenated is True
        assert len(result.df) == 5
        assert len(result.sheet_info) == 3

    def test_different_columns_first_sheet_only(self, tmp_dir):
        path = os.path.join(tmp_dir, "diff.xlsx")
        _create_xlsx(path, [
            ("S1", [["A", "B"], ["1", "2"]], []),
            ("S2", [["X", "Y", "Z"], ["3", "4", "5"]], []),
        ])

        result = preprocess_excel(path)
        assert result.sheets_concatenated is False
        assert len(result.df) == 1
        assert result.df.columns.tolist() == ["A", "B"]
        assert any("结构不同" in w for w in result.warnings)

    def test_multi_sheet_with_title_rows(self, tmp_dir):
        path = os.path.join(tmp_dir, "multi_title.xlsx")
        _create_xlsx(path, [
            ("S1", [
                ["报告标题", None, None],
                ["A", "B", "C"],
                ["1", "2", "3"],
            ], ["A1:C1"]),
            ("S2", [
                ["报告标题", None, None],
                ["A", "B", "C"],
                ["4", "5", "6"],
            ], ["A1:C1"]),
        ])

        result = preprocess_excel(path)
        assert result.sheets_concatenated is True
        assert len(result.df) == 2
        for info in result.sheet_info:
            assert info.title_rows_skipped == 1


# ── 向后兼容 ─────────────────────────────────────────────


class TestBackwardCompatibility:

    def test_simple_xlsx_structure(self, tmp_dir):
        path = os.path.join(tmp_dir, "simple.xlsx")
        _create_xlsx(path, [("Sheet1", [
            ["名称", "数值", "日期"],
            ["基金A", 1.05, "2026-01-01"],
            ["基金B", 2.10, "2026-01-02"],
        ], [])])

        result = preprocess_excel(path)
        assert result.sheet_info[0].title_rows_skipped == 0
        assert result.sheet_info[0].header_levels == 1
        assert result.df.columns.tolist() == ["名称", "数值", "日期"]
        assert len(result.df) == 2
        assert result.df.iloc[0]["名称"] == "基金A"


# ── 边界情况 ─────────────────────────────────────────────


class TestEdgeCases:

    def test_empty_workbook(self, tmp_dir):
        path = os.path.join(tmp_dir, "empty.xlsx")
        wb = openpyxl.Workbook()
        wb.save(path)

        result = preprocess_excel(path)
        assert len(result.df) == 0

    def test_single_row_data(self, tmp_dir):
        path = os.path.join(tmp_dir, "one_row.xlsx")
        _create_xlsx(path, [("Sheet1", [
            ["A", "B"],
            ["x", "y"],
        ], [])])

        result = preprocess_excel(path)
        assert len(result.df) == 1

    def test_unsupported_extension(self):
        with pytest.raises(ValueError, match="不支持"):
            preprocess_excel("test.txt")

    def test_to_str_conversions(self):
        assert _to_str(None) == ''
        assert _to_str(float('nan')) == ''
        assert _to_str(7.0) == '7'
        assert _to_str(3.14) == '3.14'
        assert _to_str('hello') == 'hello'
        assert _to_str('') == ''

    def test_build_column_names_dedup(self):
        names = _build_column_names([["A", "B", "A"]])
        assert names == ["A", "B", "A_1"]

    def test_build_column_names_empty_fallback(self):
        names = _build_column_names([["X", None, "Z"]])
        assert names[0] == "X"
        assert names[1] == "列2"
        assert names[2] == "Z"


# ── 样本文件验证 ─────────────────────────────────────────


SAMPLE_DIR = "/root/.claude/uploads/1749c10c-6322-5d16-abad-b7660e0b1d83"


class TestSampleFiles:

    @pytest.fixture
    def scenario_a(self):
        path = os.path.join(SAMPLE_DIR, "fd2b210a-_______.xls")
        if not os.path.exists(path):
            pytest.skip("Sample xls file not available")
        return path

    @pytest.fixture
    def scenario_b(self):
        path = os.path.join(SAMPLE_DIR, "0bb1094d-_____.xlsx")
        if not os.path.exists(path):
            pytest.skip("Sample xlsx file not available")
        return path

    @pytest.fixture
    def scenario_c(self):
        path = os.path.join(SAMPLE_DIR, "6efcde76-__________xls.xls")
        if not os.path.exists(path):
            pytest.skip("Sample xls file not available")
        return path

    def test_scenario_a_title_skip(self, scenario_a):
        result = preprocess_excel(scenario_a)
        assert result.sheet_info[0].title_rows_skipped == 1
        assert "持仓日期" in result.sheet_info[0].columns
        assert result.sheet_info[0].col_count == 61

    def test_scenario_b_double_header(self, scenario_b):
        result = preprocess_excel(scenario_b)
        assert result.sheet_info[0].header_levels == 2
        cols = result.df.columns.tolist()
        assert "资产名称" in cols
        assert "总资产_成本" in cols
        assert "总资产_市值" in cols
        assert "净值波动%_较上一日" in cols

    def test_scenario_c_multi_sheet_concat(self, scenario_c):
        result = preprocess_excel(scenario_c)
        assert result.sheets_concatenated is True
        assert len(result.sheet_info) == 3
        for info in result.sheet_info:
            assert info.title_rows_skipped == 1
        total_rows = sum(i.row_count for i in result.sheet_info)
        assert len(result.df) == total_rows
        assert "序号" in result.sheet_info[0].columns
