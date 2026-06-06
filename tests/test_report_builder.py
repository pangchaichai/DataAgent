"""
tests/test_report_builder.py — 报告生成模块单元测试（I-2）
"""

import os
import sys
import tempfile
from pathlib import Path

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))


# ═══════════════════════════════════════════════════════════════
#  list_templates
# ═══════════════════════════════════════════════════════════════

class TestListTemplates:
    def test_returns_list(self):
        from tools.report_builder import list_templates
        result = list_templates()
        assert isinstance(result, list)

    def test_includes_concentration_report(self):
        from tools.report_builder import list_templates
        names = {t["name"] for t in list_templates()}
        assert "concentration_report" in names

    def test_includes_nav_report(self):
        from tools.report_builder import list_templates
        names = {t["name"] for t in list_templates()}
        assert "nav_report" in names

    def test_each_entry_has_required_keys(self):
        from tools.report_builder import list_templates
        for t in list_templates():
            assert "name" in t
            assert "label" in t
            assert "path" in t


# ═══════════════════════════════════════════════════════════════
#  render_report
# ═══════════════════════════════════════════════════════════════

class TestRenderReport:
    def _concentration_data(self, breaches=None):
        return {
            "title": "主体集中度测试报告",
            "product_name": "测试产品A",
            "report_date": "2026-06-06",
            "data_date": "2026-06-05",
            "author": "测试人员",
            "market_value_field": "穿透后市值",
            "use_group_merge": True,
            "threshold_pct": 10.0,
            "breaches": breaches or [],
            "warnings": [],
        }

    def test_concentration_render_no_breaches(self):
        from tools.report_builder import render_report
        data = self._concentration_data()
        result = render_report("concentration_report", data)
        assert result.ok
        assert len(result.markdown) > 50
        assert "主体集中度" in result.markdown
        assert "无超标" in result.markdown

    def test_concentration_render_with_breaches(self):
        from tools.report_builder import render_report
        data = self._concentration_data(breaches=[
            {"entity_or_bond": "象屿集团", "product": "产品A",
             "concentration_pct": 12.5, "threshold_pct": 10.0,
             "market_value": 1234567.89},
        ])
        result = render_report("concentration_report", data)
        assert result.ok
        assert "象屿集团" in result.markdown
        assert "12.50" in result.markdown

    def test_concentration_render_with_warnings(self):
        from tools.report_builder import render_report
        data = self._concentration_data(breaches=[
            {"entity_or_bond": "X集团", "product": "产品B",
             "concentration_pct": -1.0, "threshold_pct": 10.0,
             "market_value": 100.0},
        ])
        data["warnings"] = ["⚠️ 集中度应在 0-100%：concentration_pct=-1.0（X集团）"]
        result = render_report("concentration_report", data)
        assert result.ok
        assert "数据异常提示" in result.markdown
        assert "集中度应在" in result.markdown

    def test_nav_render_basic(self):
        from tools.report_builder import render_report
        data = {
            "title": "净值运作报告",
            "product_name": "产品A",
            "report_date": "2026-06-06",
            "data_date": "2026-06-05",
            "author": "测试",
            "metrics": [
                {"product": "产品A", "nav_date": "2026-06-05",
                 "unit_nav": 1.0523, "return_7d": 0.05,
                 "return_1m": 0.12, "return_ytd": 3.5},
            ],
            "warnings": [],
        }
        result = render_report("nav_report", data)
        assert result.ok
        assert "产品A" in result.markdown
        assert "1.0523" in result.markdown

    def test_nonexistent_template_returns_error(self):
        from tools.report_builder import render_report
        result = render_report("nonexistent_template", {})
        assert not result.ok
        assert "不存在" in result.error

    def test_report_date_auto_injected(self):
        """未传 report_date 时应自动注入当日日期"""
        from tools.report_builder import render_report
        data = {
            "title": "测试报告",
            "breaches": [],
            "warnings": [],
            "use_group_merge": True,
            "threshold_pct": 10.0,
            "market_value_field": "穿透后市值",
        }
        result = render_report("concentration_report", data)
        assert result.ok
        from datetime import datetime
        today = datetime.now().strftime("%Y-%m-%d")
        assert today in result.markdown


# ═══════════════════════════════════════════════════════════════
#  export_word
# ═══════════════════════════════════════════════════════════════

class TestExportWord:
    def test_basic_export(self, tmp_path):
        from tools.report_builder import export_word
        md = "# 测试报告\n\n## 一、数据\n\n这是一段正文。\n\n- 项目一\n- 项目二\n"
        out = str(tmp_path / "test.docx")
        result = export_word(md, out)
        assert result.ok
        assert Path(result.word_path).exists()
        assert Path(result.word_path).stat().st_size > 0

    def test_table_export(self, tmp_path):
        from tools.report_builder import export_word
        md = "# 报告\n\n| 主体 | 集中度 | 阈值 |\n|------|--------|------|\n| 象屿集团 | 12.50% | 10% |\n"
        out = str(tmp_path / "table_test.docx")
        result = export_word(md, out)
        assert result.ok
        # 验证 Word 文件可读（用 python-docx 解析）
        from docx import Document
        doc = Document(result.word_path)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        tables_text = "\n".join(
            cell.text for table in doc.tables for row in table.rows for cell in row.cells
        )
        assert "象屿集团" in tables_text or "象屿集团" in full_text

    def test_heading_levels(self, tmp_path):
        from tools.report_builder import export_word
        from docx import Document
        md = "# 一级标题\n\n## 二级标题\n\n### 三级标题\n\n正文内容。\n"
        out = str(tmp_path / "headings.docx")
        result = export_word(md, out)
        assert result.ok
        doc = Document(result.word_path)
        styles = [p.style.name for p in doc.paragraphs if p.text]
        assert any("Heading 1" in s for s in styles)
        assert any("Heading 2" in s for s in styles)

    def test_output_dir_created(self, tmp_path):
        from tools.report_builder import export_word
        nested = str(tmp_path / "nested" / "dir" / "report.docx")
        md = "# 测试\n\n正文。\n"
        result = export_word(md, nested)
        assert result.ok
        assert Path(result.word_path).exists()

    def test_bold_text(self, tmp_path):
        from tools.report_builder import export_word
        from docx import Document
        md = "这是 **加粗文字** 测试。\n"
        out = str(tmp_path / "bold.docx")
        result = export_word(md, out)
        assert result.ok
        doc = Document(result.word_path)
        bold_runs = [
            run for p in doc.paragraphs for run in p.runs if run.bold
        ]
        assert len(bold_runs) >= 1


# ═══════════════════════════════════════════════════════════════
#  render + export 端到端
# ═══════════════════════════════════════════════════════════════

class TestRenderAndExport:
    def test_concentration_full_pipeline(self, tmp_path):
        from tools.report_builder import render_report, export_word
        data = {
            "title": "集中度报告 2026-06-06",
            "product_name": "全部产品",
            "report_date": "2026-06-06",
            "data_date": "2026-06-05",
            "author": "DataAgent",
            "market_value_field": "穿透后市值",
            "use_group_merge": True,
            "threshold_pct": 10.0,
            "breaches": [
                {"entity_or_bond": "象屿集团", "product": "产品A",
                 "concentration_pct": 11.2, "threshold_pct": 10.0,
                 "market_value": 5_000_000.0},
            ],
            "warnings": [],
        }
        render = render_report("concentration_report", data)
        assert render.ok

        out = str(tmp_path / "final.docx")
        word = export_word(render.markdown, out)
        assert word.ok
        assert Path(word.word_path).exists()
        # 文件大小合理（大于 5KB）
        assert Path(word.word_path).stat().st_size > 5000
