"""
tests/test_file_reader.py — 文档解析单元测试（I-6）
"""

import os
import sys
import tempfile

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))


# ═══════════════════════════════════════════════════════════════
#  TXT 文件读取测试
# ═══════════════════════════════════════════════════════════════

class TestReadText:
    def _write_txt(self, content: str, encoding='utf-8') -> str:
        with tempfile.NamedTemporaryFile(
            mode='w', suffix='.txt', delete=False, encoding=encoding
        ) as f:
            f.write(content)
            return f.name

    def test_read_utf8_txt(self):
        from tools.file_reader import read_document
        path = self._write_txt("这是测试文本。\n第二行内容。")
        try:
            result = read_document(path)
            assert result.ok
            assert "测试文本" in result.text
            assert result.file_type == "txt"
            assert result.word_count > 0
        finally:
            os.unlink(path)

    def test_read_nonexistent_file(self):
        from tools.file_reader import read_document
        result = read_document("/nonexistent/file.txt")
        assert not result.ok
        assert "不存在" in result.error

    def test_unsupported_format(self):
        from tools.file_reader import read_document
        with tempfile.NamedTemporaryFile(suffix='.csv', delete=False) as f:
            f.write(b"a,b,c\n1,2,3\n")
            path = f.name
        try:
            result = read_document(path)
            assert not result.ok
            assert ".csv" in result.error or "不支持" in result.error
        finally:
            os.unlink(path)

    def test_max_chars_truncation(self):
        from tools.file_reader import read_document
        long_content = "A" * 200
        path = self._write_txt(long_content)
        try:
            result = read_document(path, max_chars=50)
            assert result.ok
            assert len(result.text) <= 60  # some buffer for truncation message
        finally:
            os.unlink(path)


# ═══════════════════════════════════════════════════════════════
#  DOCX 文件读取测试
# ═══════════════════════════════════════════════════════════════

class TestReadDocx:
    def _create_docx(self, paragraphs: list[str], tables: list | None = None) -> str:
        from docx import Document
        doc = Document()
        for p in paragraphs:
            doc.add_paragraph(p)
        if tables:
            for headers, rows in tables:
                tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
                for j, h in enumerate(headers):
                    tbl.rows[0].cells[j].text = h
                for i, row in enumerate(rows):
                    for j, v in enumerate(row):
                        tbl.rows[i + 1].cells[j].text = v
        with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as f:
            doc.save(f.name)
            return f.name

    def test_read_docx_text(self):
        from tools.file_reader import read_document
        path = self._create_docx(["第一段", "第二段", "结论段落"])
        try:
            result = read_document(path)
            assert result.ok
            assert "第一段" in result.text
            assert result.file_type == "docx"
        finally:
            os.unlink(path)

    def test_read_docx_with_table(self):
        from tools.file_reader import read_document
        path = self._create_docx(
            ["报告标题"],
            tables=[
                (["产品", "市值", "比例"], [["产品A", "100", "10%"], ["产品B", "200", "20%"]])
            ]
        )
        try:
            result = read_document(path)
            assert result.ok
            assert len(result.tables) == 1
            tbl = result.tables[0]
            assert tbl.headers == ["产品", "市值", "比例"]
            assert len(tbl.rows) == 2
        finally:
            os.unlink(path)

    def test_docx_word_count(self):
        from tools.file_reader import read_document
        path = self._create_docx(["one two three four five"])
        try:
            result = read_document(path)
            assert result.ok
            assert result.word_count >= 5
        finally:
            os.unlink(path)
